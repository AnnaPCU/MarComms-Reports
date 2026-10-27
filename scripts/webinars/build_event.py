# ════════════════════════════════════════════════════════════════
# WEBINAR → métricas del reporte mixto (Teams + Mailchimp).
#
# Lee el Excel de lead scoring que arma el equipo por evento (hojas
# Dashboard · Priority Leads · Lead Scoring · Registrations Raw ·
# Attendance Raw · Q&A Raw · Scoring Model) y los exports de destinatarios
# de Mailchimp de la campaña (members_*.csv, una fila por destinatario), y
# emite un JSON con TODOS los números del evento: registros, asistencia,
# países, empresas, engagement, leads priorizados, envíos de email y la
# atribución de registros al canal email (en base / abrieron / con clic).
#
# Los textos narrativos del seed (highlight, lecturas, plan de acción) NO
# salen de acá: se redactan sobre estos números. Regla de honestidad: lo
# que el export no trae (LinkedIn, costo, encuestas) queda en null/[].
#
# Uso:
#   python3 scripts/webinars/build_event.py "<leads.xlsx>" <carpeta con members_*.csv> [out.json]
# Requiere: pip install openpyxl
# ════════════════════════════════════════════════════════════════
import csv, glob, json, os, re, statistics, sys, warnings, collections
import openpyxl

warnings.filterwarnings('ignore')

# Países como los escribe Teams (idioma del navegador del registrado) → nombre en español.
COUNTRY_ES = {
    'magyarország': 'Hungría', 'hungary': 'Hungría', '台灣': 'Taiwán', 'taiwan': 'Taiwán', 'ישראל': 'Israel',
    'deutschland': 'Alemania', 'germany': 'Alemania', 'suiza': 'Suiza', 'switzerland': 'Suiza', 'schweiz': 'Suiza',
    'österreich': 'Austria', 'austria': 'Austria', 'italia': 'Italia', 'italy': 'Italia', 'españa': 'España', 'spain': 'España',
    'netherlands': 'Países Bajos', 'nederland': 'Países Bajos', 'belgium': 'Bélgica', 'belgië': 'Bélgica', 'belgique': 'Bélgica',
    'france': 'Francia', 'portugal': 'Portugal', 'sweden': 'Suecia', 'sverige': 'Suecia', 'united states': 'Estados Unidos',
    'united kingdom': 'Reino Unido', 'pakistan': 'Pakistán', 'sri lanka': 'Sri Lanka', 'india': 'India', 'bulgaria': 'Bulgaria',
    'romania': 'Rumania', 'românia': 'Rumania', 'lithuania': 'Lituania', 'lietuva': 'Lituania', 'serbia': 'Serbia', 'thailand': 'Tailandia',
    'guatemala': 'Guatemala', 'chile': 'Chile', 'el salvador': 'El Salvador', 'brazil': 'Brasil', 'brasil': 'Brasil', 'mexico': 'México', 'méxico': 'México',
    'argentina': 'Argentina', 'colombia': 'Colombia', 'peru': 'Perú', 'perú': 'Perú', 'ecuador': 'Ecuador', 'poland': 'Polonia', 'polska': 'Polonia',
    'turkey': 'Turquía', 'türkiye': 'Turquía', 'south africa': 'Sudáfrica', 'greece': 'Grecia', 'ελλάδα': 'Grecia', 'denmark': 'Dinamarca',
    'danmark': 'Dinamarca', 'finland': 'Finlandia', 'suomi': 'Finlandia', 'norway': 'Noruega', 'norge': 'Noruega', 'ireland': 'Irlanda',
    'czechia': 'Chequia', 'česko': 'Chequia', 'slovakia': 'Eslovaquia', 'slovenia': 'Eslovenia', 'croatia': 'Croacia', 'hrvatska': 'Croacia',
    'egypt': 'Egipto', 'morocco': 'Marruecos', 'vietnam': 'Vietnam', 'việt nam': 'Vietnam', 'indonesia': 'Indonesia', 'malaysia': 'Malasia',
    'philippines': 'Filipinas', 'china': 'China', '中国': 'China', 'japan': 'Japón', '日本': 'Japón', 'korea': 'Corea del Sur',
    'united arab emirates': 'Emiratos Árabes Unidos', 'saudi arabia': 'Arabia Saudita', 'canada': 'Canadá', 'australia': 'Australia',
    'bangladesh': 'Bangladesh', 'nepal': 'Nepal', 'uruguay': 'Uruguay', 'paraguay': 'Paraguay', 'bolivia': 'Bolivia', 'costa rica': 'Costa Rica',
    'honduras': 'Honduras', 'nicaragua': 'Nicaragua', 'panama': 'Panamá', 'panamá': 'Panamá', 'dominican republic': 'República Dominicana',
    'ukraine': 'Ucrania', 'україна': 'Ucrania', 'latvia': 'Letonia', 'estonia': 'Estonia', 'luxembourg': 'Luxemburgo', 'cyprus': 'Chipre',
    'malta': 'Malta', 'kenya': 'Kenia', 'nigeria': 'Nigeria', 'ghana': 'Ghana', 'tunisia': 'Túnez', 'israel': 'Israel', 'iran': 'Irán',
    'singapore': 'Singapur', 'hong kong': 'Hong Kong', 'new zealand': 'Nueva Zelanda', 'russia': 'Rusia', 'россия': 'Rusia',
}


def country(v):
    s = str(v or '').strip()
    if not s:
        return None
    return COUNTRY_ES.get(s.lower(), s)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def sheet(wb, name):
    rows = [r for r in wb[name].iter_rows(values_only=True)]
    # El Excel de EmpCo (sep 2026) llama «Internal / Partner» a la columna que en Plastic era «Internal».
    hdr = [{'Internal / Partner': 'Internal'}.get(str(c).strip(), str(c).strip()) if c is not None else '' for c in rows[0]]
    return [dict(zip(hdr, r)) for r in rows[1:] if any(c is not None for c in r)]


def dur_label(minutes):
    m = int(round(minutes))
    return f'{m // 60} h {m % 60:02d} min' if m >= 60 else f'{m} min'


def parse_teams_duration(s):
    """'1 h 43 min 53s' → minutos (float)."""
    h = re.search(r'(\d+)\s*h', s or ''); m = re.search(r'(\d+)\s*min', s or ''); sec = re.search(r'(\d+)\s*s\b', s or '')
    return (int(h.group(1)) * 60 if h else 0) + (int(m.group(1)) if m else 0) + (int(sec.group(1)) / 60 if sec else 0)


def load_csv(path):
    with open(path, encoding='utf-8-sig') as fh:
        return list(csv.DictReader(fh))


# Países declarados que no son el nombre del país (formulario libre de Teams).
COUNTRY_FIX = {'usa': 'Estados Unidos', 'us': 'Estados Unidos', 'u.s.': 'Estados Unidos', 'new york ,ny': 'Estados Unidos',
               'canada ontario': 'Canadá'}
# Ambiguos: no se asignan a un país (no se deduce). «USA/CA» puede ser EE.UU./Canadá o California.
COUNTRY_AMBIGUOUS = {'usa/ca'}


def country2(v):
    s = re.sub(r'\s+', ' ', str(v or '')).strip()
    if not s or s.lower() in COUNTRY_AMBIGUOUS:
        return None
    return COUNTRY_FIX.get(s.lower()) or country(s)


# Envíos posteriores al evento (a registrados / participantes): suman como
# envíos, pero NO cuentan para la atribución de registros al canal email —
# le llegaron a la gente justamente por haberse registrado.
POST_EVENT = re.compile(r'participants|registrations|registrados|asistentes|post', re.I)


def email_block(folder, regs, registered):
    """Envíos de Mailchimp (members_*.csv) y atribución de los registrados
    (solo con los envíos previos al evento)."""
    sends = []
    by_contact = collections.defaultdict(lambda: [0, 0])
    pre = collections.defaultdict(lambda: [0, 0])
    for f in sorted(glob.glob(os.path.join(folder, 'members_*.csv'))):
        rows_csv = load_csv(f)
        sent = len(rows_csv)
        op = sum(1 for r in rows_csv if int(r.get('opens') or 0) > 0)
        cl = sum(1 for r in rows_csv if int(r.get('clicks') or 0) > 0)
        post = bool(POST_EVENT.search(os.path.basename(f)))
        sends.append({'file': os.path.basename(f), 'sent': sent, 'opens': op, 'clicks': cl, 'post': post,
                      'open': round(op / sent * 100, 1) if sent else 0, 'click': round(cl / sent * 100, 1) if sent else 0})
        for r in rows_csv:
            e = r['email'].strip().lower(); o = int(r.get('opens') or 0); c = int(r.get('clicks') or 0)
            by_contact[e][0] += o; by_contact[e][1] += c
            if not post:
                pre[e][0] += o; pre[e][1] += c
    in_base = regs & set(pre)
    email = {
        'sends': sends,
        'totalSent': sum(s['sent'] for s in sends),
        'uniqueContacts': len(by_contact),
        'openedOnce': sum(1 for v in by_contact.values() if v[0] > 0),
        'clickedOnce': sum(1 for v in by_contact.values() if v[1] > 0),
        'regInBase': len(in_base),
        'regOpened': sum(1 for e in in_base if pre[e][0] > 0),
        'regFromEmail': sum(1 for e in in_base if pre[e][1] > 0),
        'regOutsideEmail': len(regs - set(pre)),
        'attributionSends': sum(1 for x in sends if not x['post']),
    }
    email['openedOncePct'] = round(email['openedOnce'] / email['uniqueContacts'] * 100, 1)
    email['clickedOncePct'] = round(email['clickedOnce'] / email['uniqueContacts'] * 100, 1)
    email['regInBasePct'] = round(email['regInBase'] / registered * 100)
    email['regFromEmailPct'] = round(email['regFromEmail'] / registered * 100)
    email['regOutsideEmailPct'] = round(email['regOutsideEmail'] / registered * 100)
    return email


def main_segments(wb, folder):
    """Excel «por segmentos» (Fueling Conversations, oct 2026): hojas Resumen ·
    HubSpot · Base externa · Excluidos · Registro original · Asistencia original.
    El equipo ya depuró duplicados (un registro por persona) y separó internos
    y speakers; no hay score numérico sino segmentos (Hot potencial · Warm ·
    Revisar encaje · Cold · Nurture). Todo se recalcula desde «Base externa»
    (una fila por registrado externo) y «Excluidos» (internos y speakers)."""
    B = [x for x in sheet(wb, 'Base externa') if x.get('Email')]
    X = [x for x in sheet(wb, 'Excluidos') if x.get('Email')]
    R = [x for x in sheet(wb, 'Registro original') if x.get('Correo electrónico de registro')]
    A = [x for x in sheet(wb, 'Asistencia original') if x.get('Nombre')]
    att = [x for x in B if x['Webinar attendance'] == 'Attended']
    internal = [x for x in X if x['Motivo'] == 'Interno']
    speakers = [x for x in X if x['Motivo'] == 'Speaker']
    int_reg = [x for x in internal if x['Registrado'] == 'Sí']
    int_att = [x for x in internal if x['Asistió'] == 'Sí']
    registered = len(B) + len(int_reg)
    attended = len(att) + len(int_att) + sum(1 for x in speakers if x['Asistió'] == 'Sí')

    # Países declarados por los externos (registrados y asistentes).
    reg_c = collections.Counter(country2(x['Country/Region']) for x in B if country2(x['Country/Region']))
    att_c = collections.Counter(country2(x['Country/Region']) for x in att if country2(x['Country/Region']))
    countries = [{'name': c, 'reg': n, 'att': att_c.get(c, 0)} for c, n in reg_c.most_common(10)]

    # Empresas de los asistentes externos (sin empresa: dominio del email).
    GENERIC = ('gmail.', 'hotmail.', 'yahoo.', 'outlook.', 'live.', 'icloud.', 'protonmail.', 'msn.')
    def dom(x):
        d = str(x['Email'] or '').strip().lower().split('@')[-1]
        return None if not d or d.startswith(GENERIC) else d
    def org(x):
        o = re.sub(r'\s+', ' ', str(x['Company Name'] or '')).strip()
        return o or dom(x) or '—'
    prio = [x for x in att if x['Webinar segment'] in ('Hot potencial', 'Warm')]
    featured, seen = [], set()
    for x in sorted(prio, key=lambda x: -num(x['Attendance minutes'])):
        k = dom(x) or org(x)
        if k not in seen:
            seen.add(k); featured.append(org(x))
    others = []
    for x in sorted(att, key=lambda x: -num(x['Attendance minutes'])):
        k = dom(x) or org(x)
        if k not in seen:
            seen.add(k); others.append(org(x))
    unique_companies = len({dom(x) or org(x) for x in att})

    mins = [num(x['Attendance minutes']) for x in att]
    seg = collections.Counter(x['Webinar segment'] for x in B)
    organizer = [x for x in A if x['Rol'] == 'Organizador']
    room = max((parse_teams_duration(x['Duración de la reunión']) for x in organizer), default=None)
    rows = [{
        'empresa': org(x),
        'pais': country2(x['Country/Region']) or '—',
        'det': f"{num(x['Attendance minutes']):.1f}".replace('.', ',') + ' min',
        'score': None,
        'tier': 'HOT' if x['Webinar segment'] == 'Hot potencial' else 'WARM',
        'segment': x['Webinar segment'],
        'priority': x['Follow-up priority'],
    } for x in sorted(prio, key=lambda x: (str(x['Follow-up priority']), -num(x['Attendance minutes'])))]
    # Atribución con el email de registro (como en los otros eventos).
    regs = {x['Email'].strip().lower() for x in B} | {x['Email'].strip().lower() for x in int_reg}
    out = {
        'format': 'segments',
        'title': next((str(r[0]) for r in wb['Resumen'].iter_rows(values_only=True) if r and r[0]), ''),
        'registered': registered,
        'attended': attended,
        'showRateBasis': 'external',
        'showRate': round(len(att) / len(B) * 100, 1),
        'regCountries': len(reg_c),
        'countries': countries,
        'countryUnknown': sum(1 for x in B if not country2(x['Country/Region'])),
        'internos': {'total': len(int_reg), 'attended': len(int_att), 'attendedUnregistered': sum(1 for x in int_att if x['Registrado'] != 'Sí')},
        'speakers': sum(1 for x in speakers if x['Asistió'] == 'Sí'),
        'externos': {'registered': len(B), 'attended': len(att)},
        'companies': {'uniqueExternalAttended': unique_companies, 'featured': featured, 'others': others},
        'deals': {'total': len(prio), 'hot': seg.get('Hot potencial', 0), 'warm': seg.get('Warm', 0),
                  'review': seg.get('Revisar encaje', 0), 'cold': seg.get('Cold', 0), 'nurture': seg.get('Nurture', 0)},
        'durationRoomMin': round(room, 1) if room else None,
        'durationRoomLabel': dur_label(room) if room else None,
        'durationAvgMin': round(statistics.mean(mins), 1),
        'durationMedianMin': round(statistics.median(mins), 1),
        'engagement': {'high': sum(m >= 45 for m in mins), 'mid': sum(20 <= m < 45 for m in mins), 'low': sum(m < 20 for m in mins), 'base': len(att)},
        'questionsExternal': sum(int(num(x.get('Interacción: preguntas de Preguntas y respuestas'))) for x in A
                                 if str(x.get('Correo electrónico de registro') or '').lower() in {e for e in regs}),
        'hotLeads': rows,
        'email': email_block(folder, regs, registered),
        'registrationWindow': [R[0]['Hora de registro'], R[-1]['Hora de registro']] if R else None,
    }
    return out


def main():
    xlsx, folder = sys.argv[1], sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) > 3 else None
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if 'Base externa' in wb.sheetnames:
        out = main_segments(wb, folder)
        js = json.dumps(out, ensure_ascii=False, indent=1)
        if out_path:
            open(out_path, 'w').write(js); print('wrote', out_path, file=sys.stderr)
        else:
            print(js)
        d, e = out['deals'], out['email']
        print(f"  registrados {out['registered']} ({out['externos']['registered']} externos) · asistentes {out['attended']} "
              f"({out['externos']['attended']} externos, {out['showRate']}% sobre externos) · hot {d['hot']} · warm {d['warm']} · "
              f"revisar {d['review']} · cold {d['cold']} · nurture {d['nurture']}", file=sys.stderr)
        print(f"  emails: {len(e['sends'])} envíos · {e['totalSent']} enviados · {e['uniqueContacts']} contactos · reg en base {e['regInBase']} · "
              f"abrieron {e['regOpened']} · clic {e['regFromEmail']} · fuera {e['regOutsideEmail']}", file=sys.stderr)
        return

    L = [x for x in sheet(wb, 'Lead Scoring') if x.get('Email')]
    R = [x for x in sheet(wb, 'Registrations Raw') if x.get('Correo electrónico de registro')]
    P = [x for x in sheet(wb, 'Priority Leads') if x.get('Email')]
    A = [x for x in sheet(wb, 'Attendance Raw') if x.get('Nombre')]
    QA = [x for x in sheet(wb, 'Q&A Raw') if x.get('Type') == 'QUESTION']
    model = {str(r[0]).strip(): r[1:] for r in wb['Scoring Model'].iter_rows(values_only=True) if r and r[0]}

    internal = {x['Email'].lower() for x in L if x['Internal'] == 'Yes'}
    attended = [x for x in L if x['Attended'] == 'Yes']
    ext_att = [x for x in attended if x['Internal'] == 'No']
    registered = len(L)

    # ── Países: registrados (top 10 externos) y asistentes ──
    reg_c = collections.Counter(country(x['Country/Region']) for x in L if x['Internal'] == 'No' and country(x['Country/Region']))
    att_c = collections.Counter(country(x['Country/Region']) for x in ext_att if country(x['Country/Region']))
    countries = [{'name': c, 'reg': n, 'att': att_c.get(c, 0)} for c, n in reg_c.most_common(10)]
    reg_countries_all = collections.Counter(country(x['País o región']) for x in R if country(x['País o región']))

    # ── Empresas ──
    norm_org = lambda o: re.sub(r'\s+', ' ', str(o or '')).strip()
    prio_orgs = []
    for x in P:
        o = norm_org(x['Organization'])
        if o and o not in prio_orgs:
            prio_orgs.append(o)
    att_orgs = []
    for x in ext_att:
        o = norm_org(x['Organization'])
        if o and o not in att_orgs:
            att_orgs.append(o)
    others = [o for o in att_orgs if o not in prio_orgs]

    # ── Duración y engagement (asistentes, sobre el target del modelo) ──
    durs = [num(x['Duration (min)']) for x in attended]
    pcts = [num(x['Attendance %']) for x in attended]
    organizer = next((x for x in A if x['Rol'] == 'Organizador'), None)
    total_min = parse_teams_duration(organizer['Duración de la reunión']) if organizer else None
    target = num(model.get('Target attendance minutes', [None])[0]) or None

    # ── Leads priorizados ──
    def tier(s):
        return 'HOT' if str(s).lower() == 'hot' else 'WARM'
    # Sin «Organization» en el Excel (EmpCo: 53 de 94 priorizados), se muestra el
    # dominio del email corporativo — dato real del export, no una deducción.
    # Los proveedores genéricos (gmail, hotmail…) quedan en «—».
    GENERIC = ('gmail.', 'hotmail.', 'yahoo.', 'outlook.', 'live.', 'icloud.', 'protonmail.', 'msn.')
    def org_or_domain(x):
        o = norm_org(x['Organization'])
        if o:
            return o
        dom = str(x['Email'] or '').strip().lower().split('@')[-1]
        return '—' if not dom or dom.startswith(GENERIC) else dom
    rows = [{
        'empresa': org_or_domain(x),
        'pais': country(x['Country/Region']) or '—',
        'det': f"{num(x['Duration (min)']):.1f}".replace('.', ',') + ' min',
        'score': round(num(x['Score']), 1),
        'tier': tier(x['Lead Status']),
        'questions': int(num(x['Questions'])),
    } for x in P]
    hot = sum(1 for r in rows if r['tier'] == 'HOT'); warm = len(rows) - hot
    status = collections.Counter(x['Lead Status'] for x in L if x['Internal'] == 'No')

    # ── Q&A ──
    qa = [{'email': x['Identity'], 'internal': str(x['Identity']).lower() in internal, 'text': (x['Content'] or '').strip()} for x in QA]

    # ── Email (Mailchimp members: una fila por destinatario) ──
    sends = []
    by_contact = collections.defaultdict(lambda: [0, 0])
    for f in sorted(glob.glob(os.path.join(folder, 'members_*.csv'))):
        rows_csv = load_csv(f)
        sent = len(rows_csv)
        op = sum(1 for r in rows_csv if int(r.get('opens') or 0) > 0)
        cl = sum(1 for r in rows_csv if int(r.get('clicks') or 0) > 0)
        sends.append({'file': os.path.basename(f), 'sent': sent, 'opens': op, 'clicks': cl,
                      'open': round(op / sent * 100, 1) if sent else 0, 'click': round(cl / sent * 100, 1) if sent else 0})
        for r in rows_csv:
            e = r['email'].strip().lower(); by_contact[e][0] += int(r.get('opens') or 0); by_contact[e][1] += int(r.get('clicks') or 0)
    regs = {x['Email'].strip().lower() for x in L}
    in_base = regs & set(by_contact)
    email = {
        'sends': sends,
        'totalSent': sum(s['sent'] for s in sends),
        'uniqueContacts': len(by_contact),
        'openedOnce': sum(1 for v in by_contact.values() if v[0] > 0),
        'clickedOnce': sum(1 for v in by_contact.values() if v[1] > 0),
        'regInBase': len(in_base),
        'regOpened': sum(1 for e in in_base if by_contact[e][0] > 0),
        'regFromEmail': sum(1 for e in in_base if by_contact[e][1] > 0),
        'regOutsideEmail': len(regs - set(by_contact)),
    }
    email['openedOncePct'] = round(email['openedOnce'] / email['uniqueContacts'] * 100, 1)
    email['clickedOncePct'] = round(email['clickedOnce'] / email['uniqueContacts'] * 100, 1)
    email['regInBasePct'] = round(email['regInBase'] / registered * 100)
    email['regFromEmailPct'] = round(email['regFromEmail'] / registered * 100)
    email['regOutsideEmailPct'] = round(email['regOutsideEmail'] / registered * 100)

    out = {
        'title': next((str(r[0]) for r in wb['Dashboard'].iter_rows(values_only=True) if r and r[0]), ''),
        'registered': registered,
        'attended': len(attended),
        'showRate': round(len(attended) / registered * 100, 1),
        'regCountries': len(reg_countries_all),
        'countries': countries,
        'internos': {'total': len(internal), 'attended': len(attended) - len(ext_att)},
        'externos': {'registered': registered - len(internal), 'attended': len(ext_att)},
        'companies': {'uniqueExternalAttended': len(att_orgs), 'featured': prio_orgs, 'others': others},
        'deals': {'total': len(rows), 'hot': hot, 'warm': warm, 'cold': status.get('Cold', 0)},
        'durationTotalMin': round(total_min, 1) if total_min else None,
        'durationTotalLabel': dur_label(total_min) if total_min else None,
        'durationAvgMin': round(statistics.mean(durs), 1),
        'durationMedianMin': round(statistics.median(durs), 1),
        'durationTarget': target,
        'engagement': {'high': sum(p >= 0.8 for p in pcts), 'mid': sum(0.5 <= p < 0.8 for p in pcts), 'low': sum(p < 0.5 for p in pcts)},
        'qa': qa,
        'scoringModel': {k: [c for c in v if c is not None] for k, v in model.items()},
        'hotLeads': rows,
        'email': email,
        'registrationWindow': [R[0]['Hora de registro'], R[-1]['Hora de registro']] if R else None,
    }
    js = json.dumps(out, ensure_ascii=False, indent=1)
    if out_path:
        open(out_path, 'w').write(js); print('wrote', out_path, file=sys.stderr)
    else:
        print(js)
    e = out['email']
    print(f"  registrados {registered} · asistentes {len(attended)} ({out['showRate']}%) · externos {len(ext_att)} · hot {hot} · warm {warm} · cold {out['deals']['cold']}", file=sys.stderr)
    print(f"  emails: {len(sends)} envíos · {e['totalSent']} enviados · {e['uniqueContacts']} contactos · reg en base {e['regInBase']} · abrieron {e['regOpened']} · clic {e['regFromEmail']} · fuera {e['regOutsideEmail']}", file=sys.stderr)


if __name__ == '__main__':
    main()
